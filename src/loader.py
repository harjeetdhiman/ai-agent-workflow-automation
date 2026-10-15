import os
import json
from typing import Dict, List, Optional
import openpyxl
from src.schemas import WorkflowDefinition, TestQuestion, StepDefinition
from src.config import EXCEL_WORKFLOW_FILE


class ExcelWorkflowLoader:
    """
    Authoritative parser for loading workflow definitions, steps, and evaluation questions
    directly from the company's Excel workbook.
    """

    REQUIRED_WORKFLOW_COLUMNS = [
        "Workflow_ID",
        "Workflow_Name",
        "Trigger",
        "Inputs",
        "Steps",
        "Decision_Logic",
        "Tools_Required",
        "Expected_Output",
    ]

    REQUIRED_QUESTION_COLUMNS = [
        "Workflow_ID",
        "Test_Request",
        "What_To_Check",
    ]

    def __init__(self, excel_path: Optional[str] = None):
        self.excel_path = excel_path or EXCEL_WORKFLOW_FILE
        self._workflows: Dict[str, WorkflowDefinition] = {}
        self._test_questions: List[TestQuestion] = []
        self.load()

    def load(self) -> None:
        """Loads and validates worksheets from the Excel file."""
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(
                f"Excel workflow file not found at '{self.excel_path}'. "
                f"Please ensure the file is present in the specified location."
            )

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)

        if "Workflows" not in wb.sheetnames:
            raise ValueError(f"Workbook is missing the mandatory 'Workflows' sheet. Found: {wb.sheetnames}")

        self._load_workflows_sheet(wb["Workflows"])

        if "Workflow_Steps" in wb.sheetnames:
            self._load_workflow_steps_sheet(wb["Workflow_Steps"])

        # For workflows without explicit step rows, synthesize them dynamically from the Steps column
        self._synthesize_missing_steps()

        if "Test_Questions" in wb.sheetnames:
            self._load_test_questions_sheet(wb["Test_Questions"])

    def _load_workflows_sheet(self, ws) -> None:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            raise ValueError("The 'Workflows' worksheet is empty or missing data rows.")

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        missing_cols = [col for col in self.REQUIRED_WORKFLOW_COLUMNS if col not in headers]
        if missing_cols:
            raise ValueError(f"The 'Workflows' sheet is missing required columns: {missing_cols}")

        col_idx = {col: headers.index(col) for col in self.REQUIRED_WORKFLOW_COLUMNS}
        self._workflows.clear()

        for row_num, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue

            wf_id = str(row[col_idx["Workflow_ID"]]).strip() if row[col_idx["Workflow_ID"]] else ""
            if not wf_id:
                continue

            wf_name = str(row[col_idx["Workflow_Name"]]).strip() if row[col_idx["Workflow_Name"]] else f"Unnamed Workflow {wf_id}"
            trigger = str(row[col_idx["Trigger"]]).strip() if row[col_idx["Trigger"]] else ""
            inputs = str(row[col_idx["Inputs"]]).strip() if row[col_idx["Inputs"]] else ""
            steps_raw = str(row[col_idx["Steps"]]).strip() if row[col_idx["Steps"]] else ""
            decision_logic = str(row[col_idx["Decision_Logic"]]).strip() if row[col_idx["Decision_Logic"]] else ""
            tools_raw = str(row[col_idx["Tools_Required"]]).strip() if row[col_idx["Tools_Required"]] else ""
            expected_output = str(row[col_idx["Expected_Output"]]).strip() if row[col_idx["Expected_Output"]] else ""

            arrow_sep = "→" if "→" in steps_raw else ("->" if "->" in steps_raw else None)
            if arrow_sep:
                steps = [s.strip() for s in steps_raw.split(arrow_sep) if s.strip()]
            else:
                steps = [steps_raw] if steps_raw else []

            tools = [t.strip() for t in tools_raw.replace(";", ",").split(",") if t.strip()]

            wf_def = WorkflowDefinition(
                workflow_id=wf_id,
                workflow_name=wf_name,
                trigger=trigger,
                inputs=inputs,
                steps_raw=steps_raw,
                steps=steps,
                decision_logic=decision_logic,
                tools_required=tools,
                expected_output=expected_output,
                executable_steps=[],
            )
            self._workflows[wf_id] = wf_def

    def _load_workflow_steps_sheet(self, ws) -> None:
        """Loads declarative step definitions from the Workflow_Steps worksheet."""
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        h_map = {h.lower(): idx for idx, h in enumerate(headers)}

        for row in rows[1:]:
            if not any(row):
                continue
            wf_id_idx = h_map.get("workflow_id", 0)
            wf_id = str(row[wf_id_idx]).strip() if row[wf_id_idx] else ""
            if not wf_id or wf_id not in self._workflows:
                continue

            order_idx = h_map.get("step_order", 1)
            name_idx = h_map.get("step_name", 2)
            tool_idx = h_map.get("tool_name", 3)
            params_idx = h_map.get("params", 4)
            source_idx = h_map.get("input_source", 5)
            cond_idx = h_map.get("condition_expr", h_map.get("condition", 6))
            on_true_idx = h_map.get("on_true", None)
            on_false_idx = h_map.get("on_false", None)
            err_idx = h_map.get("on_error", 7 if on_true_idx is None else 9)

            try:
                step_order = int(row[order_idx]) if row[order_idx] is not None else len(self._workflows[wf_id].executable_steps) + 1
            except (ValueError, TypeError):
                step_order = len(self._workflows[wf_id].executable_steps) + 1

            step_name = str(row[name_idx]).strip() if row[name_idx] else f"Step {step_order}"
            tool_name = str(row[tool_idx]).strip() if row[tool_idx] else "llm_generate"

            raw_params = row[params_idx] if len(row) > params_idx else "{}"
            params = {}
            if raw_params:
                if isinstance(raw_params, dict):
                    params = raw_params
                elif isinstance(raw_params, str) and raw_params.strip():
                    try:
                        params = json.loads(raw_params)
                    except json.JSONDecodeError:
                        params = {"raw_params": raw_params}

            input_source = str(row[source_idx]).strip() if len(row) > source_idx and row[source_idx] else None
            cond_val = str(row[cond_idx]).strip() if cond_idx is not None and len(row) > cond_idx and row[cond_idx] else None
            condition_expr = cond_val if cond_val and cond_val.lower() not in ["none", "null", ""] else None
            on_true = str(row[on_true_idx]).strip().lower() if on_true_idx is not None and len(row) > on_true_idx and row[on_true_idx] else "continue"
            on_false = str(row[on_false_idx]).strip().lower() if on_false_idx is not None and len(row) > on_false_idx and row[on_false_idx] else "continue"
            on_error = str(row[err_idx]).strip().lower() if err_idx is not None and len(row) > err_idx and row[err_idx] else "stop"

            step_def = StepDefinition(
                step_number=step_order,
                step_name=step_name,
                tool_name=tool_name,
                params=params,
                input_source=input_source,
                condition=condition_expr,
                condition_expr=condition_expr,
                on_true=on_true,
                on_false=on_false,
                on_error=on_error,
                output_variable=f"step_{step_order}.output",
                summary_template=f"Executed {step_name} via {tool_name}"
            )
            self._workflows[wf_id].executable_steps.append(step_def)

    def _synthesize_missing_steps(self) -> None:
        """Synthesizes step definitions dynamically for workflows missing in the Workflow_Steps sheet."""
        for wf_id, wf in self._workflows.items():
            if not wf.executable_steps and wf.steps:
                synthesized = []
                for idx, step_name in enumerate(wf.steps, start=1):
                    tool = wf.tools_required[idx - 1] if idx <= len(wf.tools_required) else (wf.tools_required[-1] if wf.tools_required else "llm_generate")
                    source = f"step_{idx-1}.output" if idx > 1 else "params"
                    synthesized.append(
                        StepDefinition(
                            step_number=idx,
                            step_name=step_name,
                            tool_name=tool,
                            params={},
                            input_source=source,
                            condition=None,
                            on_error="continue",
                            output_variable=f"step_{idx}.output",
                            summary_template=f"Executed {step_name} via {tool}"
                        )
                    )
                wf.executable_steps = synthesized

    def _load_test_questions_sheet(self, ws) -> None:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        col_idx = {col: headers.index(col) for col in self.REQUIRED_QUESTION_COLUMNS if col in headers}
        self._test_questions.clear()

        for row in rows[1:]:
            if not any(row):
                continue
            wf_id = str(row[col_idx.get("Workflow_ID", 0)]).strip() if row[col_idx.get("Workflow_ID", 0)] else ""
            test_req = str(row[col_idx.get("Test_Request", 1)]).strip() if row[col_idx.get("Test_Request", 1)] else ""
            check = str(row[col_idx.get("What_To_Check", 2)]).strip() if row[col_idx.get("What_To_Check", 2)] else ""

            if wf_id and test_req:
                self._test_questions.append(
                    TestQuestion(
                        workflow_id=wf_id,
                        test_request=test_req,
                        what_to_check=check
                    )
                )
                if wf_id in self._workflows:
                    self._workflows[wf_id].test_examples.append(test_req)

    def get_workflow(self, workflow_id: str) -> Optional[WorkflowDefinition]:
        wf = self._workflows.get(workflow_id)
        if wf and not wf.executable_steps and wf.steps:
            self._synthesize_missing_steps()
        return wf

    def get_all_workflows(self) -> Dict[str, WorkflowDefinition]:
        return dict(self._workflows)

    def get_test_questions(self) -> List[TestQuestion]:
        return list(self._test_questions)


# Global singleton instance
workflow_loader = ExcelWorkflowLoader()
WorkflowLoader = ExcelWorkflowLoader
